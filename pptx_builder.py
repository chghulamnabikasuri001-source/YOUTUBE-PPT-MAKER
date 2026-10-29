"""Builds the .pptx file (themes + screenshots)."""
import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

# dark = title slide background, accent = top bar, light = content background,
# text = bullet colour, sub = title slide subtitle
THEMES = {
    "Navy & Teal": dict(dark="1F2A44", accent="0F9D8A", light="F5F7FA", text="222B3A", sub="BFE9E2", muted="8892A6"),
    "Forest Green": dict(dark="14532D", accent="16A34A", light="F3FAF5", text="1F2937", sub="BBF7D0", muted="7A8B80"),
    "Maroon": dict(dark="7F1D1D", accent="DC2626", light="FBF5F5", text="1F2937", sub="FECACA", muted="8F7A7A"),
    "Purple": dict(dark="3B0764", accent="9333EA", light="F8F5FC", text="1F2937", sub="E9D5FF", muted="877A99"),
    "Charcoal & Orange": dict(dark="111827", accent="F97316", light="F9FAFB", text="1F2937", sub="FED7AA", muted="8A8F99"),
}
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _c(hex_str):
    return RGBColor.from_string(hex_str)


def _rtl(p, rtl):
    if rtl:
        p._p.get_or_add_pPr().set("rtl", "1")
        p.alignment = PP_ALIGN.RIGHT


def _text(slide, x, y, w, h, text, size, color, bold=False, rtl=False, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = "Arial"
    r.font.color.rgb = color
    _rtl(p, rtl)
    return tb


def _bg(slide, hex_str):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = _c(hex_str)


def _bar(slide, prs, hex_str, height):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = _c(hex_str)
    shp.line.fill.background()


def _content_slide(prs, T, heading, bullets, notes, rtl, label="", image=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, T["light"])
    _bar(s, prs, T["accent"], Inches(0.18))
    _text(s, Inches(0.7), Inches(0.5), Inches(11.9), Inches(1.0), heading, 32,
          _c(T["dark"]), bold=True, rtl=rtl, anchor=MSO_ANCHOR.MIDDLE)

    total = sum(len(str(b)) for b in bullets)
    if image:
        size = 22 if total < 350 else 18 if total < 600 else 15
        text_w = Inches(6.5)
        if rtl:
            img_x, text_x = Inches(0.7), Inches(6.3)
        else:
            img_x, text_x = Inches(7.2), Inches(0.7)
        try:
            s.shapes.add_picture(io.BytesIO(image), img_x, Inches(1.9), width=Inches(5.4))
        except Exception:
            pass  # if the image is corrupt, the slide is built without it
    else:
        size = 24 if total < 450 else 20 if total < 750 else 17
        text_w = Inches(11.5)
        text_x = Inches(0.9)

    tb = s.shapes.add_textbox(text_x, Inches(1.8), text_w, Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = "•  " + str(b)
        r.font.size = Pt(size)
        r.font.name = "Arial"
        r.font.color.rgb = _c(T["text"])
        p.space_after = Pt(12)
        _rtl(p, rtl)

    if label:
        _text(s, Inches(11.6), Inches(6.95), Inches(1.4), Inches(0.4), label, 11, _c(T["muted"]))
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def build_pptx(title, channel, video_url, slides, summary, quiz, rtl, theme_name="Navy & Teal") -> bytes:
    """slides: [{"heading","bullets","notes","image"(optional bytes)}]"""
    T = THEMES.get(theme_name, THEMES["Navy & Teal"])
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Title slide
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, T["dark"])
    _bar(s, prs, T["accent"], Inches(0.3))
    _text(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(2.0), title, 42, WHITE,
          bold=True, rtl=rtl, anchor=MSO_ANCHOR.MIDDLE)
    credit = f"Source: {channel}" if channel else "Source: YouTube"
    _text(s, Inches(0.9), Inches(4.6), Inches(11.5), Inches(0.6), credit, 20, _c(T["sub"]))
    _text(s, Inches(0.9), Inches(5.2), Inches(11.5), Inches(0.5), video_url, 14, _c(T["muted"]))

    n = len(slides)
    for i, sl in enumerate(slides, 1):
        _content_slide(prs, T, sl.get("heading", ""), sl.get("bullets", []), sl.get("notes", ""),
                       rtl, f"{i} / {n}", sl.get("image"))

    if summary:
        _content_slide(prs, T, "خلاصہ" if rtl else "Summary", summary, "", rtl)
    if quiz:
        qs = [f"Q{i}. {q.get('q', '')}" for i, q in enumerate(quiz, 1)]
        ans = [f"Q{i}. {q.get('a', '')}" for i, q in enumerate(quiz, 1)]
        _content_slide(prs, T, "کوئز" if rtl else "Quiz", qs, "", rtl)
        _content_slide(prs, T, "جوابات" if rtl else "Answers", ans, "", rtl)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
